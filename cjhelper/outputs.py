"""Write the generated files: one CSV per table, scenarios.csv, report.md, testbook.xlsx and the zip.

Spec, "Data generation rules", file format: "UTF-8, comma-separated, header row with the exact field
names in schema order, True/False for Booleans, empty cell for null."

Spec, "Outputs": "The app returns one zip file with the data CSVs plus three supporting files":
``<TableName>.csv`` (one per selected table), ``scenarios.csv``, ``testbook.xlsx`` and ``report.md``.

Every function here is pure: the same input gives the same bytes ("The output is reproducible").
For that reason the timestamps that Excel files and zip archives normally carry are fixed.
"""

from __future__ import annotations

import csv
import dataclasses
import io
import re
import zipfile
from datetime import datetime, timezone
from typing import Iterable, Optional

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from .datagen import SCENARIO_COLUMNS, activation_setting, boolean_text, case_paths, wait_precedes_manual_action
from .model import (
    KIND_WAIT_ATTRIBUTE,
    KIND_WAIT_DURATION,
    KIND_WAIT_STO,
    LEVEL_SCENARIO,
    STATUS_COVERED,
    STATUS_MANUAL,
    STATUS_NOT_COVERED,
    Case,
    GenerationResult,
    Journey,
    Package,
    Scenario,
    Settings,
    Step,
    Table,
    TableData,
    UnsupportedFilter,
    Value,
    Write,
)

CSV_LINE_END = "\n"
SCENARIOS_FILE = "scenarios.csv"
REPORT_FILE = "report.md"
TESTBOOK_FILE = "testbook.xlsx"

# Spec, "Columns of testbook.xlsx".
TESTBOOK_COLUMNS = ["ID", "Owner system", "Test name", "Description", "Pre conditions",
                    "Detailed steps", "Expected result", "Status", "Note"]
OWNER_SYSTEM = "SFMC"
STATUS_TO_START = "TO START"
TESTBOOK_SHEET = "Testbook"
_TESTBOOK_WIDTHS = [6, 12, 40, 60, 60, 70, 70, 12, 40]

ZIP_DATE_TIME = (1980, 1, 1, 0, 0, 0)    # fixed, so the same files always give the same zip bytes
_CORE_TIMESTAMP = re.compile(r"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)")
_CURRENT_TIME_VALUE = "{{DateTime.Now}}"
_REPORT_DETAIL_LIMIT = 50
SUSPICIOUS_PREFIX = "suspicious pattern:"   # how other modules mark a suspicious pattern


def table_file_name(table_name: str) -> str:
    """Spec, "Outputs": one ``<TableName>.csv`` per selected table."""
    return f"{table_name}.csv"


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def schema_field_names(table: Table) -> list[str]:
    """The exact field names in schema order."""
    return [f.name for f in sorted(table.fields, key=lambda f: f.ordinal)]


def _csv_bytes(header: list[str], rows: Iterable[list[Value]]) -> bytes:
    """UTF-8 without byte order mark, comma separated, minimal quoting, "\\n" line ends, None -> empty cell."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator=CSV_LINE_END)
    writer.writerow(header)
    for values in rows:
        writer.writerow(["" if value is None else value for value in values])
    return buffer.getvalue().encode("utf-8")


def _row_value(row: dict[str, Value], column: str) -> Value:
    if column in row:
        return row[column]
    low = column.lower()
    for key, value in row.items():
        if key.lower() == low:
            return value
    return None


def write_table_csv(data: TableData, table: Table) -> bytes:
    """One data CSV: the header is the table's schema, in schema order, whatever the rows hold."""
    header = schema_field_names(table)
    return _csv_bytes(header, ([_row_value(row, column) for column in header] for row in data.rows))


def write_scenarios_csv(rows: list[dict[str, Value]]) -> bytes:
    """scenarios.csv: one row per scenario; the columns are those of the rows, in first-seen order."""
    header: list[str] = []
    for row in rows:
        header.extend(column for column in row if column not in header)
    if not header:
        header = list(SCENARIO_COLUMNS)
    return _csv_bytes(header, ([row.get(column) for column in header] for row in rows))


def read_csv(data: bytes) -> tuple[list[str], list[dict[str, Value]]]:
    """Read a CSV written here back: (header, rows). An empty cell reads as None.

    Strict UTF-8, so a byte order mark would show up in the first column name. Blank lines are
    skipped, cells beyond the header are dropped and missing cells read as None.
    """
    reader = csv.reader(io.StringIO(data.decode("utf-8"), newline=""))
    lines = list(reader)
    if not lines:
        return [], []
    header = lines[0]
    rows = [{column: (line[i] if i < len(line) and line[i] != "" else None) for i, column in enumerate(header)}
            for line in lines[1:] if line]
    return header, rows


# ---------------------------------------------------------------------------
# Zip
# ---------------------------------------------------------------------------

def _zip_info(name: str) -> zipfile.ZipInfo:
    info = zipfile.ZipInfo(name, date_time=ZIP_DATE_TIME)
    info.compress_type = zipfile.ZIP_DEFLATED
    info.external_attr = 0o644 << 16
    return info


def build_zip(files: dict[str, bytes]) -> bytes:
    """Spec, "Outputs": "one zip file with the data CSVs plus three supporting files". Files keep the
    given order; every entry has the same fixed timestamp."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(_zip_info(name), content)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Shared wording helpers
# ---------------------------------------------------------------------------

def _show(value: Value) -> str:
    return "(empty)" if value is None or value == "" else str(value)


def _plural(count: int, singular: str, plural: str) -> str:
    return f"{count} {singular if count == 1 else plural}"


def _verb(count: int) -> str:
    return "is" if count == 1 else "are"


def _field_type(package: Package, table_name: str, field_name: str) -> str:
    table = next((t for t in package.tables.values() if t.name.lower() == table_name.lower()), None)
    field = table.field(field_name) if table is not None else None
    return field.type.lower() if field is not None and field.type else ""


def _show_value(package: Package, table_name: str, field_name: str, value: Value) -> str:
    """A value as the CSV holds it: Booleans as True/False."""
    if _field_type(package, table_name, field_name) == "boolean":
        value = boolean_text(value) or value
    return _show(value)


def _duration_text(amount, unit: Optional[str]) -> str:
    unit_text = (unit or "").lower()
    if amount == 1 and unit_text.endswith("s"):
        unit_text = unit_text[:-1]
    return f"{amount if amount is not None else '?'} {unit_text}".strip()


def _date(moment: datetime, settings: Settings) -> str:
    return moment.strftime(settings.date_format)


def _activity_name(journey: Journey, key: Optional[str]) -> str:
    """An activity by its name, or its key when it has none."""
    if key is None:
        return "(none)"
    activity = journey.activities.get(key)
    return activity.name if activity is not None and activity.name else key


def _activity_label(journey: Journey, key: Optional[str]) -> str:
    """``Name (KEY)``, or the key alone when the activity has no name."""
    name = _activity_name(journey, key)
    return name if name == key else f"{name} ({key})"


def _step_text(journey: Journey, step: Step) -> str:
    """An activity of the path; for an activity with several outcomes, the outcome taken."""
    text = _activity_name(journey, step.activity_key)
    activity = journey.activities.get(step.activity_key)
    if activity is not None and len(activity.outcomes) > 1 and step.outcome_key is not None:
        outcome = activity.outcome(step.outcome_key)
        text += f" [{outcome.label if outcome is not None and outcome.label else step.outcome_key}]"
    return text


def _path_text(journey: Journey, steps: Iterable[Step]) -> str:
    texts = [_step_text(journey, step) for step in steps]
    return " > ".join(texts) if texts else "(no activity)"


def _email_names(journey: Journey, keys: Iterable[str]) -> str:
    names = []
    for key in keys:
        activity = journey.activities.get(key)
        names.append((activity.email_name or activity.name or key) if activity is not None else key)
    return ", ".join(names) if names else "none"


def _writes_text(writes: Iterable[Write]) -> str:
    """The final value of each field written by the journey, in the order first written."""
    final: dict[tuple[str, str], Value] = {}
    for write in writes:
        final[(write.table, write.field)] = write.value
    parts = []
    for (table, field), value in final.items():
        shown = "the current date and time" if value == _CURRENT_TIME_VALUE else _show(value)
        parts.append(f"{table}.{field} = {shown}")
    return "; ".join(parts) if parts else "none"


def _ordered_scenarios(result: GenerationResult) -> list[Scenario]:
    return sorted(result.coverage.scenarios, key=lambda s: (s.run_order, s.id))


def _key_range(keys: list[str]) -> str:
    if not keys:
        return "none"
    if len(keys) <= 3:
        return ", ".join(keys)
    return f"{keys[0]} to {keys[-1]}"


def _smallest_branch(case: Case):
    branches = case.random.branches if case.random is not None else []
    return min(branches, key=lambda b: b.percentage) if branches else None


def _outcome_label(journey: Journey, activity_key: str, outcome_key: str) -> str:
    activity = journey.activities.get(activity_key)
    outcome = activity.outcome(outcome_key) if activity is not None else None
    return outcome.label if outcome is not None and outcome.label else outcome_key


def _branch_label(journey: Journey, random_key: str, branch) -> str:
    """A random split branch by its label, with its percentage when the label does not show it."""
    label = _outcome_label(journey, random_key, branch.outcome_key)
    return label if f"{branch.percentage}%" in label else f"{label} ({branch.percentage}%)"


def _path_keys(case: Case) -> list[str]:
    """Every activity a case can cross, in path order: the deterministic part, then the random branches."""
    keys: list[str] = []
    for path in case_paths(case):
        for key in path:
            if key not in keys:
                keys.append(key)
    return keys


# ---------------------------------------------------------------------------
# Testbook
# ---------------------------------------------------------------------------

@dataclasses.dataclass
class _TestbookContext:
    result: GenerationResult
    journey: Journey
    settings: Settings
    csv_names: list[str]
    one_active: bool
    activation_ref: Optional[tuple[str, str]]
    first_scenario: Optional[Scenario]
    emails: dict[str, str]           # contact key -> address


def _testbook_context(result: GenerationResult) -> _TestbookContext:
    on, ref, _ = activation_setting(result.package, result.table_model, result.settings)
    scenarios = _ordered_scenarios(result)
    return _TestbookContext(
        result=result,
        journey=result.package.journey,
        settings=result.settings,
        csv_names=[table_file_name(name) for name in result.data.tables],
        one_active=on,
        activation_ref=ref,
        first_scenario=scenarios[0] if scenarios else None,
        emails={contact.key: contact.email for contact in result.data.contacts},
    )


def build_testbook_rows(result: GenerationResult) -> list[list]:
    """Spec, "Columns of testbook.xlsx": one row per case, scenarios in run order."""
    context = _testbook_context(result)
    rows: list[list] = []
    for scenario in _ordered_scenarios(result):
        for case in scenario.cases:
            rows.append([
                len(rows) + 1,
                OWNER_SYSTEM,
                _test_name(context, case),
                _description(context, scenario, case),
                _preconditions(context, scenario, case),
                _detailed_steps(context, scenario, case),
                _expected_result(context, case),
                STATUS_TO_START,
                _note(context, scenario, case),
            ])
    return rows


def _test_name(context: _TestbookContext, case: Case) -> str:
    """Spec: "Case identifier and a short summary of the path"."""
    summary = case.summary or f"Path ending at {_activity_name(context.journey, case.end_activity_key)}"
    return f"{case.id} - {summary}"


def _values_text(package: Package, values: dict, order: list) -> str:
    keys = [ref for ref in order if ref in values] + [ref for ref in values if ref not in order]
    return "; ".join(f"{table}.{field} = {_show_value(package, table, field, values[(table, field)])}"
                     for table, field in keys)


def _description(context: _TestbookContext, scenario: Scenario, case: Case) -> str:
    """Spec: "Scenario values, contact profile values, and the expected path as activity names in order"."""
    coverage = context.result.coverage
    package = context.result.package
    scenario_values = _values_text(package, scenario.values, coverage.scenario_fields) or "no scenario-level value"
    contact_values = _values_text(package, case.contact_values, coverage.contact_fields) or "no contact-level value"
    lines = [
        f"Scenario {scenario.id} ({scenario.key}): {scenario_values}.",
        f"Contact profile: {contact_values}.",
        f"Expected path: {_path_text(context.journey, case.steps)}.",
    ]
    if case.random is not None:
        branches = [f"{_branch_label(context.journey, case.random.activity_key, b)}: "
                    f"{_path_text(context.journey, b.steps) if b.steps else 'leaves the journey'}"
                    for b in case.random.branches]
        lines.append("After the random split, each contact takes one branch: " + "; ".join(branches) + ".")
    return "\n".join(lines)


def _scenario_key_column(context: _TestbookContext, table_name: str) -> Optional[str]:
    """The field of a table that holds the scenario key."""
    for group in context.result.table_model.key_groups:
        if group.kind != "scenario":
            continue
        for table, field in group.fields:
            if table.lower() == table_name.lower():
                return field
    return None


def _activation_text(context: _TestbookContext, scenario: Scenario) -> str:
    """Spec: "When on, only the first scenario is active in the CSV, and the testbook tells the tester
    which row to switch before each run"; "To move to the next scenario, the tester edits only its
    activation field"."""
    table, field = context.activation_ref
    key_column = _scenario_key_column(context, table)
    row_text = f"the row with {key_column} {scenario.key}" if key_column else f"the row of scenario {scenario.key}"
    if scenario is context.first_scenario:
        return (f"Check that scenario {scenario.key} is the active one: in {table}, {field} is True on {row_text} "
                "and False on every other row, as generated.")
    return (f"When the contacts of the previous scenario have left the journey, activate scenario {scenario.key}: "
            f"in {table}, set {field} to True on {row_text} and to False on every other row. Change nothing else.")


def _preconditions(context: _TestbookContext, scenario: Scenario, case: Case) -> str:
    """Spec: "Which CSVs must be imported, which scenario must be active, and the keys of the test
    contacts"; "The 50 contacts of a random split case are given as a key range"."""
    lines = [f"Import these CSV files: {', '.join(context.csv_names)}."]
    if context.one_active:
        table, field = context.activation_ref
        lines.append(f"Active scenario: {scenario.key} ({table}.{field} = True on its row only).")
    else:
        lines.append(f"Scenario: {scenario.key}. Every scenario is active at the same time.")
    keys = case.contact_keys
    if case.random is not None and len(keys) > 3:
        lines.append(f"Test contacts ({len(keys)}): {keys[0]} to {keys[-1]}.")
    else:
        lines.append(f"Test contacts ({len(keys)}): {', '.join(keys) if keys else 'none'}.")
    return "\n".join(lines)


def _wait_descriptions(context: _TestbookContext, scenario: Scenario, case: Case) -> list[str]:
    """Spec: "Wait by duration, send-time wait: The expected delay is noted in the testbook"."""
    texts = []
    for key in _path_keys(case):
        activity = context.journey.activities.get(key)
        if activity is None:
            continue
        name = _activity_name(context.journey, activity.key)
        if activity.kind == KIND_WAIT_DURATION:
            texts.append(f"{name} ({_duration_text(activity.wait_amount, activity.wait_unit)})")
        elif activity.kind == KIND_WAIT_STO:
            if activity.wait_amount:
                texts.append(f"{name} (send-time wait: sent at the best time within "
                             f"{_duration_text(activity.wait_amount, activity.wait_unit)}, so up to that long)")
            else:
                texts.append(f"{name} (send-time wait: the delay is not known in advance)")
        elif activity.kind == KIND_WAIT_ATTRIBUTE and activity.wait_field is not None:
            table, field = activity.wait_field
            value = _written_value(context, table, field, scenario, case)
            until = f"until {value}, the date in {table}.{field}" if value else f"until the date in {table}.{field}"
            texts.append(f"{name} ({until})")
    return texts


def _written_value(context: _TestbookContext, table_name: str, field_name: str,
                   scenario: Scenario, case: Case) -> Value:
    """The value written in the CSV for this scenario's or this case's row of a table."""
    data = next((d for name, d in context.result.data.tables.items() if name.lower() == table_name.lower()), None)
    if data is None or not data.rows:
        return None
    role = next((r for r in context.result.table_model.tables.values() if r.name.lower() == table_name.lower()), None)
    kind = "scenario" if role is not None and role.level == LEVEL_SCENARIO else "contact"
    owner = scenario.key if kind == "scenario" else (case.contact_keys[0] if case.contact_keys else None)
    key_columns = [field for group in context.result.table_model.key_groups if group.kind == kind
                   for table, field in group.fields if table.lower() == table_name.lower()]
    for row in data.rows:
        if any(_row_value(row, column) == owner for column in key_columns):
            return _row_value(row, field_name)
    return None


def _deadline(context: _TestbookContext, scenario: Scenario, case: Case, action) -> Optional[str]:
    """Spec: "the testbook gives the deadline as a date and time". The deadline is the date of the last
    wait by attribute between the email and the engagement split: when it passes, the split is evaluated."""
    for keys in case_paths(case):
        if action.activity_key not in keys:
            continue
        end = keys.index(action.activity_key)
        start = keys.index(action.email_activity_key) if action.email_activity_key in keys[:end] else -1
        for key in reversed(keys[start + 1:end]):
            activity = context.journey.activities.get(key)
            if (activity is not None and activity.kind == KIND_WAIT_ATTRIBUTE and activity.wait_field is not None
                    and wait_precedes_manual_action(key, case)):
                return _written_value(context, activity.wait_field[0], activity.wait_field[1], scenario, case)
    return None


def _manual_steps(context: _TestbookContext, scenario: Scenario, case: Case) -> list[str]:
    """Spec: "any manual action with its deadline"; "Contacts that must open or click get the tester's
    real inbox with plus-addressing"."""
    steps = []
    inbox = context.settings.real_inbox.strip()
    addresses = ", ".join(context.emails.get(key, key) for key in case.contact_keys)
    for action in case.manual_actions:
        verb = "open the email" if action.kind == "open" else "click a link in the email"
        deadline = _deadline(context, scenario, case, action)
        when = f"Before {deadline}" if deadline else "As soon as the email arrives"
        if not inbox:
            steps.append(f"The split {_activity_label(context.journey, action.activity_key)} needs the contacts to "
                         f"{verb} '{action.email_name}', but no real inbox was set, so the emails cannot be received.")
            continue
        steps.append(f"{when}, {verb} '{action.email_name}' in the inbox {inbox}, for each contact of this test "
                     f"(addresses {addresses}). Do nothing on the emails of other tests.")
    return steps


def _detailed_steps(context: _TestbookContext, scenario: Scenario, case: Case) -> str:
    """Spec: "Numbered steps: activate the scenario, start the entry, wait (with expected delays), any
    manual action with its deadline, collect evidence"."""
    steps = []
    if context.one_active:
        steps.append(_activation_text(context, scenario))
    else:
        steps.append(f"Check that the rows of scenario {scenario.key} are imported.")
    entry = context.result.package.entry_event
    start = f"Start the entry of journey '{context.journey.name}'"
    if entry is not None:
        start += f" (entry event '{entry.name}'"
        start += f", automation '{entry.automation_name}')" if entry.automation_name else ")"
    steps.append(start + ". The tests of this scenario run together. The entry CSV is imported directly: "
                 "importing the source tables and running the automation is the alternative, not combined with it.")
    waits = _wait_descriptions(context, scenario, case)
    steps.append("Let the contacts pass the waits on the path: " + "; ".join(waits) + "." if waits else
                 "There is no wait on the path.")
    steps.extend(_manual_steps(context, scenario, case))
    steps.append("Collect the evidence for each contact: the emails received, the values written by the journey "
                 "and the last activity reached.")
    return "\n".join(f"{number}. {text}" for number, text in enumerate(steps, 1))


def _expected_result(context: _TestbookContext, case: Case) -> str:
    """Spec: "For each contact: the emails or messages received, the values written by the journey, and
    the last activity reached". The contacts of a case share one expected result."""
    journey = context.journey
    keys = case.contact_keys
    if case.random is None:
        return "\n".join([
            f"For each contact ({', '.join(keys) if keys else 'none'}):",
            f"- Emails received: {_email_names(journey, case.emails_sent)}.",
            f"- Values written by the journey: {_writes_text(case.writes)}.",
            f"- Last activity reached: {_activity_label(journey, case.end_activity_key)}.",
        ])
    random_key = case.random.activity_key
    lines = [
        f"For each of the {len(keys)} contacts ({_key_range(keys)}), up to the random split "
        f"{_activity_label(journey, random_key)}:",
        f"- Emails received: {_email_names(journey, case.emails_sent)}.",
        f"- Values written by the journey: {_writes_text(case.writes)}.",
        "Then each contact takes one branch of the random split:",
    ]
    for branch in case.random.branches:
        lines.append(f"- {_branch_label(journey, random_key, branch)}: emails received: "
                     f"{_email_names(journey, branch.emails_sent)}; "
                     f"values written: {_writes_text(branch.writes)}; last activity reached: "
                     f"{_activity_label(journey, branch.end_activity_key)}.")
    smallest = _smallest_branch(case)
    if smallest is not None:
        lines.append(f"Accepted range: {case.random.expected_low} to {case.random.expected_high} contacts on the "
                     f"{_branch_label(journey, random_key, smallest)} side.")
    return "\n".join(lines)


def _note(context: _TestbookContext, scenario: Scenario, case: Case) -> str:
    """Spec: "Manual engagement required, or the accepted range for a random split case"."""
    notes = []
    zones = sorted({a.wait_timezone for a in context.journey.activities.values()
                    if a.kind == KIND_WAIT_ATTRIBUTE and a.wait_timezone})
    for action in case.manual_actions:
        verb = "open" if action.kind == "open" else "click"
        deadline = _deadline(context, scenario, case, action)
        text = f"Manual engagement required: {verb} the email '{action.email_name}'"
        notes.append(text + (f" before {deadline}." if deadline else " as soon as it arrives."))
        if deadline and zones:
            notes.append(f"The wait date is read in {', '.join(zones)}, so the deadline could shift by some hours.")
    if case.random is not None:
        smallest = _smallest_branch(case)
        label = _branch_label(context.journey, case.random.activity_key, smallest) if smallest else "smaller"
        notes.append(f"Random split: the contacts cannot be forced down one branch. Between {case.random.expected_low} "
                     f"and {case.random.expected_high} of the {len(case.contact_keys)} contacts are expected on the "
                     f"{label} side.")
    return " ".join(notes)


def _excel_text(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _utc_text(moment: datetime) -> str:
    if moment.tzinfo is not None:
        moment = moment.astimezone(timezone.utc)
    return moment.strftime("%Y-%m-%dT%H:%M:%SZ")


def _stable_xlsx(content: bytes, moment: datetime) -> bytes:
    """openpyxl stamps the save time into docProps/core.xml and the zip entries: replace both, so the
    same input always gives the same bytes."""
    stamp = _utc_text(moment)
    source = zipfile.ZipFile(io.BytesIO(content))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = _CORE_TIMESTAMP.sub(lambda m: m.group(1) + stamp + m.group(2), data.decode("utf-8"))
                data = text.encode("utf-8")
            archive.writestr(_zip_info(info.filename), data)
    return buffer.getvalue()


def write_testbook(result: GenerationResult) -> bytes:
    """Spec, "Outputs": "testbook.xlsx: One test per case with the expected result, in the team's
    testbook layout"; "The testbook has one row per case, not per contact"."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TESTBOOK_SHEET
    sheet.append(TESTBOOK_COLUMNS)
    for row in build_testbook_rows(result):
        sheet.append([_excel_text(value) for value in row])
    header_fill = PatternFill("solid", fgColor="1F3864")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for position, width in enumerate(_TESTBOOK_WIDTHS, 1):
        sheet.column_dimensions[sheet.cell(1, position).column_letter].width = width
    sheet.freeze_panes = "A2"
    workbook.properties.creator = "CJ Test Helper"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return _stable_xlsx(buffer.getvalue(), result.generation_time)


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------

def _md_cell(value) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\r", " ").replace("\n", " ")


def _md_table(header: list[str], rows: list[list]) -> str:
    lines = ["| " + " | ".join(_md_cell(h) for h in header) + " |",
             "| " + " | ".join("---" for _ in header) + " |"]
    lines += ["| " + " | ".join(_md_cell(value) for value in row) + " |" for row in rows]
    return "\n".join(lines)


def _bullets(items: Iterable[str]) -> str:
    return "\n".join(f"- {item}" for item in items)


def _unique(items: Iterable[str]) -> list[str]:
    seen: list[str] = []
    for item in items:
        if item not in seen:
            seen.append(item)
    return seen


def _split_suspicious(warnings: Iterable[str]) -> tuple[list[str], list[str]]:
    """Spec, "Outputs": the report "lists suspicious patterns found while parsing, such as the same value
    written in two letter cases". Other modules mark them with the prefix ``suspicious pattern:``."""
    patterns, others = [], []
    for warning in warnings:
        if warning.lower().startswith(SUSPICIOUS_PREFIX):
            pattern = warning[len(SUSPICIOUS_PREFIX):].strip()
            patterns.append(pattern[:1].upper() + pattern[1:])
        else:
            others.append(warning)
    return patterns, others


def _section_title(result: GenerationResult) -> str:
    settings = result.settings
    return "\n\n".join([
        f"# Test data report: {result.package.journey.name}",
        f"Journey {result.package.journey.name} from the package {result.package.name}, generated on "
        f"{_date(result.generation_time, settings)} with the run identifier {settings.run_id}.",
    ])


def _section_summary(result: GenerationResult) -> str:
    coverage, data = result.coverage, result.data
    cases = coverage.cases
    statuses = [b.status for b in coverage.branches]
    covered = statuses.count(STATUS_COVERED)
    manual = statuses.count(STATUS_MANUAL)
    not_covered = statuses.count(STATUS_NOT_COVERED)
    statistical = sum(1 for b in coverage.branches if b.statistical and b.status != STATUS_NOT_COVERED)
    rows = [
        ["Scenarios", len(coverage.scenarios)],
        ["Statistical scenarios", sum(1 for s in coverage.scenarios if s.is_statistical)],
        ["Cases", len(cases)],
        ["Cases with a manual action", sum(1 for c in cases if c.manual_actions)],
        ["Statistical cases (random split)", sum(1 for c in cases if c.random is not None)],
        ["Contacts", len(data.contacts)],
        ["Contacts using the real inbox", sum(1 for c in data.contacts if c.uses_real_inbox)],
        ["Branches", len(coverage.branches)],
        ["Branches covered", covered],
        ["Branches covered with manual action", manual],
        ["Branches not covered", not_covered],
        ["Branches covered only through a random split", statistical],
    ]
    text = (f"The test data has {_plural(len(coverage.scenarios), 'scenario', 'scenarios')}, "
            f"{_plural(len(cases), 'case', 'cases')} and {_plural(len(data.contacts), 'contact', 'contacts')}. "
            f"Of {_plural(len(coverage.branches), 'branch', 'branches')}, {covered} {_verb(covered)} covered, "
            f"{manual} {_verb(manual)} covered with a manual action and {not_covered} {_verb(not_covered)} not covered.")
    return "\n\n".join(["## Summary", text, _md_table(["Item", "Count"], rows)])


def _section_checks(result: GenerationResult) -> str:
    if not result.checks:
        return "## Trust checks\n\nNo trust check was run on these files."
    failed = [check for check in result.checks if not check.passed]
    text = f"{len(result.checks) - len(failed)} of the {len(result.checks)} trust checks passed."
    if failed:
        text += f" {len(failed)} failed: the details are below."
    rows = [[check.name, "passed" if check.passed else "FAILED", len(check.details)] for check in result.checks]
    parts = ["## Trust checks", text, _md_table(["Check", "Result", "Details"], rows)]
    for check in result.checks:
        if not check.details:
            continue
        shown = check.details[:_REPORT_DETAIL_LIMIT]
        lines = [f"### {check.name}", _bullets(shown)]
        if len(check.details) > len(shown):
            lines.append(f"... and {len(check.details) - len(shown)} more.")
        parts.append("\n\n".join(lines))
    return "\n\n".join(parts)


def _section_tables(result: GenerationResult) -> str:
    rows = []
    for role in result.table_model.tables.values():
        data = result.data.tables.get(role.name)
        rows.append([role.name, ", ".join(sorted(role.roles)) or "none", role.level or "-",
                     table_file_name(role.name) if role.selected else "no",
                     len(data.rows) if data is not None else "-", role.reason])
    selected = [table_file_name(name) for name in result.data.tables]
    text = f"{_plural(len(selected), 'table gets', 'tables get')} a CSV: {', '.join(selected) if selected else 'none'}."
    return "\n\n".join(["## Tables", text, _md_table(["Table", "Role", "Level", "CSV", "Rows", "Reason"], rows)])


def _section_scenarios(result: GenerationResult) -> str:
    journey = result.package.journey
    on, ref, _ = activation_setting(result.package, result.table_model, result.settings)
    intro = ("Run the scenarios in run order. Import every CSV once. Do not change or re-import the rows of a "
             "scenario until its contacts have left the journey.")
    if on:
        intro += (f" Only the first scenario is active in the CSV; to move to the next one, change only "
                  f"{ref[0]}.{ref[1]}.")
    scenario_rows = []
    for row in result.data.scenarios_rows:
        values = "; ".join(f"{column} = {_show(value)}" for column, value in row.items()
                           if column not in SCENARIO_COLUMNS)
        scenario_rows.append([row.get("scenario_key"), row.get("run_order"), row.get("activate"),
                              row.get("has_manual_action"), row.get("is_statistical"), row.get("case_count"),
                              row.get("contact_count"),
                              f"{row.get('first_contact_key') or '-'} to {row.get('last_contact_key') or '-'}",
                              values or "none"])
    case_rows = []
    for scenario in _ordered_scenarios(result):
        for case in scenario.cases:
            if case.random is not None:
                smallest = _smallest_branch(case)
                label = _branch_label(journey, case.random.activity_key, smallest) if smallest else "?"
                special = (f"random split {case.random.activity_key}: {case.random.expected_low} to "
                           f"{case.random.expected_high} contacts expected on the {label} side")
            else:
                special = "; ".join(f"{a.kind} '{a.email_name}' at {_activity_label(journey, a.activity_key)}"
                                    for a in case.manual_actions)
            case_rows.append([case.id, scenario.key, len(case.contact_keys), _key_range(case.contact_keys),
                              special or "none", _activity_label(journey, case.end_activity_key),
                              case.summary or "-"])
    return "\n\n".join([
        "## Scenarios", intro,
        _md_table(["Scenario", "Run order", "Active at start", "Manual action", "Statistical", "Cases", "Contacts",
                   "Contact keys", "Values"], scenario_rows),
        "### Cases",
        _md_table(["Case", "Scenario", "Contacts", "Contact keys", "Manual action or random split",
                   "Last activity", "Summary"], case_rows),
    ])


def _section_not_covered(result: GenerationResult) -> str:
    rows = [[f"{b.activity_name or b.activity_key} ({b.activity_key})", f"{b.outcome_label} ({b.outcome_key})",
             b.reason or "no reason given"]
            for b in result.coverage.branches if b.status == STATUS_NOT_COVERED]
    if not rows:
        return "## Branches not covered\n\nEvery branch is covered."
    text = f"{_plural(len(rows), 'branch is', 'branches are')} not covered."
    return "\n\n".join(["## Branches not covered", text, _md_table(["Activity", "Outcome", "Reason"], rows)])


def _section_unsupported(result: GenerationResult) -> str:
    journey = result.package.journey
    items = []
    for key in _unique(list(result.package.unsupported_activities) + list(result.coverage.unsupported)):
        activity = journey.activities.get(key)
        kind = f" (type {activity.raw_type})" if activity is not None else ""
        items.append(f"Activity {_activity_label(journey, key)}{kind} is not supported; the branches behind it are "
                     "not covered.")
    query = result.table_model.entry_query
    if query is not None:
        for flt in query.filters:
            if isinstance(flt, UnsupportedFilter):
                items.append(f"The entry query filter `{flt.text}` is not supported.")
    items.extend(f"Placeholder: {text}" for text in result.data.placeholders)
    if not items:
        return "## Unsupported elements\n\nNo unsupported element was found."
    return "\n\n".join(["## Unsupported elements", _bullets(items)])


def _section_warnings(result: GenerationResult) -> str:
    query = result.table_model.entry_query
    sources = [
        ("While reading the export", result.package.warnings),
        ("Table selection and SQL", list(result.table_model.warnings) + (list(query.warnings) if query else [])),
        ("Coverage", result.coverage.warnings),
        ("Data generation", result.data.warnings),
    ]
    patterns: list[str] = []
    groups: list[tuple[str, list[str]]] = []
    for title, warnings in sources:
        found, others = _split_suspicious(warnings)
        patterns.extend(found)
        groups.append((title, others))
    parts = ["## Warnings and suspicious patterns"]
    for title, items in [("Suspicious patterns", patterns)] + groups:
        items = _unique(items)
        if items:
            parts.append(f"### {title}\n\n{_bullets(items)}")
    if len(parts) == 1:
        parts.append("No warning.")
    return "\n\n".join(parts)


def _section_assumptions(result: GenerationResult) -> str:
    """Spec, "Assumptions to verify and known limits": the six assumptions the generator encodes."""
    settings = result.settings
    case_text = ("treat `Email` and `EMAIL` as the same value (setting on)" if settings.case_insensitive_text
                 else "treat `Email` and `EMAIL` as different values (setting off)")
    zones = sorted({a.wait_timezone for a in result.package.journey.activities.values()
                    if a.kind == KIND_WAIT_ATTRIBUTE and a.wait_timezone})
    zone_text = f" The waits of this journey use {', '.join(zones)}." if zones else ""
    items = [
        "Decision split outcomes are evaluated in order and the first match wins.",
        f"Text comparisons {case_text}.",
        "A wait by attribute whose date is already past ends immediately.",
        f"The import accepts dates written as {_date(result.generation_time, settings)} "
        f"(format `{settings.date_format}`) and `True`/`False` for Booleans.",
        "An engagement split is evaluated when the contact reaches it, not when the email was sent.",
        "The time zone applied to a wait date is not known, so the deadlines in the testbook could shift by some "
        f"hours.{zone_text}",
    ]
    return "\n\n".join([
        "## Assumptions",
        "The expected results depend on these assumptions about how Salesforce behaves. Each should be checked once "
        "with a small real run.",
        _bullets(items),
    ])


_SETTING_LABELS = {
    "contacts_per_case": "Contacts per case",
    "contacts_per_random_case": "Contacts per random-split case",
    "real_inbox": "Real inbox for open and click tests",
    "real_phone": "Real phone number",
    "fake_phone": "Fake phone number",
    "window_start_days": "Validity window start (days from the generation time)",
    "window_end_days": "Validity window end (days from the generation time)",
    "wait_offset_hours": "Wait date offset (hours)",
    "wait_offset_before_engagement_hours": "Wait date offset before an engagement split (hours)",
    "date_format": "Date format in the CSVs",
    "one_active_scenario": "One active scenario at a time",
    "activation_field": "Activation field",
    "run_id": "Run identifier in keys",
    "case_insensitive_text": "Text comparisons ignore letter case",
    "seed": "Seed of the set cover",
    "generation_time": "Generation time",
}


def _setting_text(name: str, value, result: GenerationResult) -> str:
    if name == "generation_time":
        return _date(result.generation_time, result.settings)
    if name == "one_active_scenario" and value is None:
        return "on when an activation field is detected"
    if value is None:
        return "(not set)"
    if isinstance(value, bool):
        return "on" if value else "off"
    if isinstance(value, tuple):
        return ".".join(str(part) for part in value)
    if value == "":
        return "(empty)"
    return str(value)


def _section_settings(result: GenerationResult) -> str:
    rows = [[_SETTING_LABELS.get(f.name, f.name), _setting_text(f.name, getattr(result.settings, f.name), result)]
            for f in dataclasses.fields(result.settings)]
    on, ref, _ = activation_setting(result.package, result.table_model, result.settings)
    rows.append(["One active scenario at a time, as applied",
                 f"on, activation field {ref[0]}.{ref[1]}" if on else "off"])
    return "\n\n".join(["## Settings used", _md_table(["Setting", "Value"], rows)])


def write_report(result: GenerationResult) -> str:
    """Spec, "Outputs": "report.md: Coverage, unreachable branches, unsupported elements, warnings,
    settings used". It "opens with a summary: number of scenarios, cases and contacts, and branches
    covered, covered with manual action, and not covered". Every number is computed from ``result``."""
    sections = [
        _section_title(result),
        _section_summary(result),
        _section_checks(result),
        _section_tables(result),
        _section_scenarios(result),
        _section_not_covered(result),
        _section_unsupported(result),
        _section_warnings(result),
        _section_assumptions(result),
        _section_settings(result),
    ]
    return "\n\n".join(sections) + "\n"
