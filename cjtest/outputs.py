"""Write a test run to disk: import CSVs, expected-results matrix, testbook, report, manifest."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .engine import OTHER, Result
from .expectations import action_text, binomial_interval, describe
from .generator import Dataset
from .model import Package
from .static_checks import Finding

TESTBOOK_HEADER = ["ID", "Owner system", "Test name", "Description", "Pre conditions",
                   "Detailed steps", "Expected result", "Status", "Note"]
STATUSES = "TO START, ON GOING, COMPLETED, TO REVIEW"


def write_run(out_dir: Path, pkg: Package, prof: dict, result: Result, ds: Dataset,
              findings: list[Finding]) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    ctx = _Context(pkg, prof, result, ds)
    files = {}
    files.update(_write_csvs(out_dir / "import", pkg, ds))
    files["matrix"] = _write_matrix(out_dir / "expected_matrix.xlsx", ctx, findings)
    files["testbook"] = _write_testbook(out_dir / "testbook.xlsx", ctx)
    files["report"] = _write_report(out_dir / "report.md", ctx, findings)
    files["manifest"] = _write_manifest(out_dir / "manifest.json", ctx)
    (out_dir / "profile.json").write_text(json.dumps(prof, indent=2, ensure_ascii=False), encoding="utf-8")
    return files


class _Context:
    def __init__(self, pkg, prof, result, ds):
        self.pkg, self.prof, self.result, self.ds = pkg, prof, result, ds
        self.s = prof["settings"]
        self.expect = {}
        for scen in result.scenarios:
            for case in scen.cases:
                self.expect[case.id] = describe(pkg, case.steps, case.engagement)
        self.contact_des = list(ds.contact_rows)
        self.scenario_des = list(ds.scenario_rows)
        self.automation = prof["entry"].get("automation")

    def fmt(self, v) -> str:
        if v is None:
            return "(empty)"
        if v == OTHER:
            return self.s["other_value"]
        return str(v)

    def assignment_text(self, assignment: dict, variables) -> str:
        return ", ".join(f"{v[1]}={self.fmt(assignment.get(v))}" for v in variables)


# -- CSV ------------------------------------------------------------------------------------
def _write_csvs(folder: Path, pkg: Package, ds: Dataset) -> dict[str, Path]:
    folder.mkdir(parents=True, exist_ok=True)
    files = {}
    for de, rows in ds.contact_rows.items():
        files[f"csv:{de}"] = _csv(folder / f"{de}.csv", pkg, de, rows)
    for de, by_scen in ds.scenario_rows.items():
        for sid, rows in by_scen.items():
            files[f"csv:{de}:{sid}"] = _csv(folder / de / f"{sid}.csv", pkg, de, rows)
    return files


def _csv(path: Path, pkg: Package, de: str, rows: list[dict]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    names = [f.name for f in pkg.des[de].fields]
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=names)
        w.writeheader()
        w.writerows(rows)
    return path


# -- Excel helpers ------------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F3864")


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: dict[int, int] | None = None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for i in range(1, len(header) + 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = (widths or {}).get(i, 18)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


# -- expected matrix ------------------------------------------------------------------------------
def _write_matrix(path: Path, ctx: _Context, findings: list[Finding]) -> Path:
    r, ds = ctx.result, ctx.ds
    wb = Workbook()
    wb.remove(wb.active)

    svars, cvars = r.scenario_vars, r.contact_vars
    scen_rows = []
    for s in r.scenarios:
        n_contacts = s.contacts if s.kind == "statistical" else len(s.cases) * ctx.s["records_per_case"]
        files = ", ".join(f"import/{de}/{s.id}.csv" for de in ctx.scenario_des)
        scen_rows.append([s.id, s.kind, ds.campaigns[s.id], *[ctx.fmt(s.assignment.get(v)) for v in svars],
                          len(s.cases), n_contacts, files])
    _sheet(wb, "Scenarios", ["Scenario", "Kind", "Campaign", *[v[1] for v in svars], "Cases", "Contacts", "Config file"],
           scen_rows, {3: 26, len(svars) + 6: 40})

    case_rows, record_rows = [], []
    for s in r.scenarios:
        for case in s.cases:
            exp = ctx.expect[case.id]
            rec = ds.case_records.get(case.id)
            keys = rec.contact_keys if rec else []
            case_rows.append([
                case.id, s.id, s.kind, *[ctx.fmt(case.contact.get(v)) for v in cvars],
                "\n".join(exp.tester_actions) or "-",
                "\n".join(exp.trail), "\n".join(action_text(a) for a in exp.actions) or "no action",
                exp.exit_activity, len(keys), "\n".join(keys) if len(keys) <= 5 else f"{keys[0]} .. {keys[-1]}"])
            for k, mail in zip(keys, rec.emails if rec else []):
                record_rows.append([k, case.id, s.id, ds.campaigns[s.id], mail, exp.summary])
    _sheet(wb, "Cases", ["Case", "Scenario", "Kind", *[v[1] for v in cvars], "Tester actions", "Decision trail",
                         "Expected actions", "Exit activity", "Records", "Contact keys"],
           case_rows, {len(cvars) + 4: 30, len(cvars) + 5: 60, len(cvars) + 6: 70, len(cvars) + 7: 28,
                       len(cvars) + 9: 34})
    _sheet(wb, "Records", ["Contact key", "Case", "Scenario", "Campaign", "Email", "Expected summary"],
           record_rows, {1: 30, 4: 26, 5: 44, 6: 60})

    acts = ctx.pkg.journey.activities
    cov_rows = []
    for g in r.goals:
        a = acts[g[0]]
        o = next(x for x in a.outcomes if x.key == g[1])
        if g in r.covered:
            status = "covered"
        elif g in r.covered_statistical:
            status = "covered (statistical scenario)"
        elif g in r.unreachable_engagement:
            status = "unreachable: referenced email never sent"
        elif g in r.blocked_by_engagement:
            status = "unreachable: below an unreachable engagement branch"
        else:
            status = "not covered"
        cov_rows.append([g[0], a.name, o.label, o.criteria.describe() if o.criteria else "(remainder)", status])
    _sheet(wb, "Coverage", ["Activity", "Name", "Outcome", "Criteria", "Status"], cov_rows, {3: 30, 4: 70, 5: 40})
    _sheet(wb, "Static checks", ["Severity", "Check", "Finding"],
           [[f.severity, f.code, f.message] for f in findings], {3: 120})
    wb.save(path)
    return path


# -- testbook --------------------------------------------------------------------------------------
def _write_testbook(path: Path, ctx: _Context) -> Path:
    r, ds, pkg = ctx.result, ctx.ds, ctx.pkg
    rows = []
    contact_files = ", ".join(f"import/{de}.csv -> {de}" for de in ctx.contact_des)
    for s in r.scenarios:
        campaign = ds.campaigns[s.id]
        cfg = ", ".join(f"import/{de}/{s.id}.csv -> {de}" for de in ctx.scenario_des)
        pre = (f"Contacts imported once for the whole run ({contact_files}).\n"
               f"Journey '{pkg.journey.name}' active.\n"
               f"No other campaign ACTIVE in {', '.join(ctx.scenario_des)}.")
        for case in s.cases:
            exp = ctx.expect[case.id]
            rec = ds.case_records[case.id]
            steps = [f"Import {cfg} (Add and Update): campaign {campaign}.",
                     f"Run once the automation '{ctx.automation}'." if ctx.automation else "Start the entry.",
                     "Wait for the contacts to complete the journey"
                     + (" (send-time optimization: up to 24h)." if any(a["type"] == "sto_wait" for a in exp.actions) else ".")]
            if exp.tester_actions and s.kind == "path":
                deadline = ds.engagement_deadlines.get(s.id, "")
                steps.append(f"In mailbox {ctx.s['engagement_mailbox']}, for each contact of this case "
                             f"(recipients {', '.join(e.split('@')[0] for e in rec.emails)}): "
                             + "; then ".join(exp.tester_actions)
                             + (f". Do it BEFORE {deadline}, when the wait before the engagement split ends"
                                if deadline else "")
                             + ". Do nothing on emails of other cases.")
            steps.append("Collect the evidence and compare it with the expected result.")
            steps_txt = "\n".join(f"{i}. {t}" for i, t in enumerate(steps, 1))
            desc = (f"Scenario {s.id}: {ctx.assignment_text(s.assignment, r.scenario_vars)}.\n"
                    f"Contact profile: {ctx.assignment_text(case.contact, r.contact_vars)}.\n"
                    f"Path: {' > '.join(exp.trail)}")
            if s.kind == "statistical":
                expected = _statistical_expected(ctx, s)
                name = f"{s.id} - Random split distribution ({s.contacts} contacts)"
                keys = f"{rec.contact_keys[0]} .. {rec.contact_keys[-1]}"
                note = "Statistical check: individual contacts are not predictable."
            else:
                acts_txt = "\n".join(f"- {action_text(a)}" for a in exp.actions) or "- no action"
                expected = (f"For each contact ({', '.join(rec.contact_keys)}):\n{acts_txt}\n"
                            f"- exits the journey after {exp.exit_activity}\n- no other action")
                name = f"{case.id} - {exp.summary}"
                keys = ", ".join(rec.contact_keys)
                note = "Requires manual email engagement." if exp.tester_actions else ""
            rows.append([len(rows) + 1, "SFMC", name, desc, pre + f"\nTest contacts: {keys}.",
                         steps_txt, expected, "TO START", note])
    wb = Workbook()
    wb.remove(wb.active)
    ws = _sheet(wb, "Testbook", TESTBOOK_HEADER, rows, {1: 6, 2: 10, 3: 40, 4: 60, 5: 50, 6: 60, 7: 70, 8: 14, 9: 30})
    dv = DataValidation(type="list", formula1=f'"{STATUSES}"', allow_blank=True)
    ws.add_data_validation(dv)
    if rows:
        dv.add(f"H2:H{len(rows) + 1}")
    wb.save(path)
    return path


def _statistical_expected(ctx: _Context, s) -> str:
    lines = [f"Across the {s.contacts} contacts:"]
    for e in s.random_expectation:
        p = (e.get("percentage") or 0) / 100
        lo, hi = binomial_interval(s.contacts, p)
        exp = describe(ctx.pkg, e["steps"], {})
        first = next((action_text(a) for a in exp.actions), "no action")
        lines.append(f"- branch '{e['label']}' ({e.get('percentage')}%): about {s.contacts * p:.0f} contacts, "
                     f"acceptable {lo}-{hi} (95%); first action: {first}")
    return "\n".join(lines)


# -- report and manifest ------------------------------------------------------------------------------
def _write_report(path: Path, ctx: _Context, findings: list[Finding]) -> Path:
    r, ds, pkg = ctx.result, ctx.ds, ctx.pkg
    path_s = [s for s in r.scenarios if s.kind == "path"]
    stat_s = [s for s in r.scenarios if s.kind == "statistical"]
    n_cases = sum(len(s.cases) for s in path_s)
    n_eng = sum(1 for s in path_s for c in s.cases if any(c.engagement.values()))
    n_contacts = sum(len(v) for v in ds.contact_rows.values()) // max(1, len(ds.contact_rows))
    lines = [
        f"# Test run {ds.run_id}: {pkg.journey.name}", "",
        "## Summary", "",
        f"- Path scenarios: {len(path_s)} (one campaign each, run in sequence)",
        f"- Statistical scenarios: {len(stat_s)}",
        f"- Test cases: {n_cases} ({n_eng} need manual email engagement)",
        f"- Contacts: {n_contacts}",
        f"- Outcomes covered: {len(r.covered)} of {len(r.goals)}"
        + (f", plus {len(r.covered_statistical)} by the statistical scenario" if r.covered_statistical else ""),
        f"- Outcomes unreachable in the journey: {len(r.unreachable_engagement) + len(r.blocked_by_engagement)}",
        f"- Other outcomes not covered: {len(r.uncovered) - len(r.unreachable_engagement) - len(r.blocked_by_engagement)}",
        "",
        "## Variables", "",
        "Scenario level (one value per campaign): " + ", ".join(f"{v[0]}.{v[1]}" for v in r.scenario_vars), "",
        "Contact level: " + ", ".join(f"{v[0]}.{v[1]}" for v in r.contact_vars), "",
        "## How to run", "",
        "1. Import once: " + "; ".join(f"`import/{de}.csv` into `{de}`" for de in ctx.contact_des) + ".",
        "2. For each scenario, in order: import `import/<DE>/<scenario>.csv` into "
        + ", ".join(f"`{de}`" for de in ctx.scenario_des)
        + (f", then run once the automation `{ctx.automation}`." if ctx.automation else ", then start the entry."),
        "3. Engagement cases: open/click the emails in the internal mailbox as described in the testbook.",
        "4. Collect evidence and fill the testbook.", "",
        "## Static checks", "",
    ]
    lines += [f"- **{f.severity}** `{f.code}`: {f.message}" for f in findings] or ["- none"]
    if ds.errors or r.warnings:
        lines += ["", "## Errors and warnings", ""]
        lines += [f"- ERROR: {e}" for e in ds.errors] + [f"- {w}" for w in r.warnings]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _write_manifest(path: Path, ctx: _Context) -> Path:
    r, ds = ctx.result, ctx.ds
    data = {
        "run_id": ds.run_id,
        "journey": ctx.pkg.journey.name,
        "journey_key": ctx.pkg.journey.key,
        "scenarios": [{"id": s.id, "kind": s.kind, "campaign": ds.campaigns[s.id],
                       "assignment": {f"{v[0]}.{v[1]}": s.assignment.get(v) for v in r.scenario_vars},
                       "contacts": s.contacts,
                       "random_expectation": [{"outcome": e["outcome"], "percentage": e.get("percentage"),
                                               "steps": [list(x) for x in e["steps"]]} for e in s.random_expectation]}
                      for s in r.scenarios],
        "cases": [{"id": c.id, "scenario": c.scenario_id,
                   "contact_keys": ds.case_records[c.id].contact_keys,
                   "emails": ds.case_records[c.id].emails,
                   "engagement": c.engagement,
                   "steps": [list(x) for x in c.steps],
                   "expected_actions": ctx.expect[c.id].actions}
                  for s in r.scenarios for c in s.cases],
    }
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    return path
