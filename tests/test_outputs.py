"""Unit tests for cjhelper.outputs (spec, "Outputs" and the file format of "Data generation rules").

The inputs come from the hand-made world of tests/test_datagen.py, run through datagen.
"""

from __future__ import annotations

import io
import zipfile

from openpyxl import load_workbook

from cjhelper.datagen import SCENARIO_COLUMNS
from cjhelper.model import CheckResult, Field, GenerationResult, Table, TableData
from cjhelper.outputs import (
    REPORT_FILE,
    SCENARIOS_FILE,
    TESTBOOK_COLUMNS,
    TESTBOOK_FILE,
    build_zip,
    read_csv,
    table_file_name,
    build_testbook_rows,
    write_report,
    write_scenarios_csv,
    write_table_csv,
    write_testbook,
)
from tests.test_datagen import CONFIG, ENTRY, MASTER, build_world, run


def build_result(checks=None, **settings) -> GenerationResult:
    """A full GenerationResult, with every file of the zip written as the pipeline would."""
    world = build_world(**settings)
    data = run(world)
    result = GenerationResult(
        package=world.package, table_model=world.model, coverage=world.coverage, data=data,
        checks=checks if checks is not None else [
            CheckResult("Values respect the schema", True, []),
            CheckResult("Every key joins", False, ["DE_Master row 4 has no scenario row.", "Second detail."]),
        ],
        settings=world.settings, files={}, generation_time=world.generation_time)
    for name, table_data in data.tables.items():
        result.files[table_file_name(name)] = write_table_csv(table_data, world.package.tables[name])
    result.files[SCENARIOS_FILE] = write_scenarios_csv(data.scenarios_rows)
    result.files[REPORT_FILE] = write_report(result).encode("utf-8")
    result.files[TESTBOOK_FILE] = write_testbook(result)
    return result


def small_table() -> Table:
    # Fields listed out of schema order on purpose: the header follows the ordinals.
    return Table(id="t", name="DE_Small", customer_key=None, fields=[
        Field("Flag", "Boolean", None, False, True, None, 2),
        Field("Id", "Text", 10, True, False, None, 0),
        Field("Name", "Text", 50, False, True, None, 1),
        Field("Note", "Text", 50, False, True, None, 3),
    ])


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def test_table_csv_format():
    data = TableData("DE_Small", ["Id", "Name", "Flag", "Note"], [
        {"Id": "1", "Name": "Caffè, \"bar\"", "Flag": "True", "Note": None},
        {"id": "2", "NAME": "plain", "flag": "False", "note": "line\nbreak"},   # other letter case: still found
    ])
    content = write_table_csv(data, small_table())
    assert not content.startswith(b"\xef\xbb\xbf")                 # no byte order mark
    assert b"\r\n" not in content.replace(b"line\nbreak", b"")     # "\n" line ends
    assert content.decode("utf-8") == (
        "Id,Name,Flag,Note\n"
        "1,\"Caffè, \"\"bar\"\"\",True,\n"
        "2,plain,False,\"line\nbreak\"\n"
    )


def test_table_csv_header_is_the_schema_even_for_missing_values():
    content = write_table_csv(TableData("DE_Small", [], [{"Id": "1"}]), small_table())
    assert content == b"Id,Name,Flag,Note\n1,,,\n"


def test_read_csv_round_trip():
    header, rows = read_csv(b"A,B,C\n1,,\"x, y\"\n\"multi\nline\",2,\n")
    assert header == ["A", "B", "C"]
    assert rows == [{"A": "1", "B": None, "C": "x, y"}, {"A": "multi\nline", "B": "2", "C": None}]
    assert read_csv(b"") == ([], [])
    assert read_csv(b"A\n\n1\n") == (["A"], [{"A": "1"}])                 # blank line skipped
    assert read_csv(write_table_csv(TableData("T", [], [{"Id": None}]), small_table()))[1] == [
        {"Id": None, "Name": None, "Flag": None, "Note": None}]


def test_generated_csvs_read_back_identical():
    result = build_result()
    for name, table_data in result.data.tables.items():
        header, rows = read_csv(result.files[table_file_name(name)])
        assert header == result.package.tables[name].field_names
        assert rows == [{column: (value if value != "" else None) for column, value in row.items()}
                        for row in table_data.rows]


def test_generated_csv_writes_booleans_and_nulls_as_the_spec_says():
    result = build_result()
    text = result.files[table_file_name(CONFIG)].decode("utf-8")
    lines = text.split("\n")
    assert lines[0] == "CampaignName,CHANNEL,ACTIVE,HAS_RUN,START_DATE,END_DATE,RUN_TIME,Region,WAIT1,Label"
    assert lines[1] == ("TST_R01_S01,Email,True,False,2026-10-10 09:30,2026-10-18 09:30,,North,"
                        "2026-10-12 09:30,TST_La")
    assert text.endswith("\n") and len(lines) == 4                 # header, 2 rows, final empty string


def test_scenarios_csv():
    result = build_result()
    header, rows = read_csv(result.files[SCENARIOS_FILE])
    assert header == SCENARIO_COLUMNS + [f"{CONFIG}.CHANNEL"]
    assert [r["scenario_key"] for r in rows] == ["TST_R01_S01", "TST_R01_S02"]
    assert [r["activate"] for r in rows] == ["yes", "no"]
    assert rows[1]["contact_count"] == "50"


def test_scenarios_csv_without_scenarios_still_has_a_header():
    assert write_scenarios_csv([]) == (",".join(SCENARIO_COLUMNS) + "\n").encode("utf-8")


# ---------------------------------------------------------------------------
# Zip
# ---------------------------------------------------------------------------

def test_zip_contains_every_file():
    result = build_result()
    archive = zipfile.ZipFile(io.BytesIO(build_zip(result.files)))
    assert archive.namelist() == list(result.files)
    assert archive.namelist() == [f"{ENTRY}.csv", f"{MASTER}.csv", f"{CONFIG}.csv", SCENARIOS_FILE, REPORT_FILE,
                                  TESTBOOK_FILE]
    for name, content in result.files.items():
        assert archive.read(name) == content


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------

def test_report_sections_in_order():
    report = write_report(build_result())
    titles = [line for line in report.splitlines() if line.startswith("## ")]
    assert titles == ["## Summary", "## Trust checks", "## Tables", "## Scenarios", "## Branches not covered",
                      "## Unsupported elements", "## Warnings and suspicious patterns", "## Assumptions",
                      "## Settings used"]


def test_report_totals_match_the_data():
    result = build_result()
    report = write_report(result)
    assert "The test data has 2 scenarios, 3 cases and 56 contacts." in report
    assert "Of 10 branches, 8 are covered, 1 is covered with a manual action and 1 is not covered." in report
    assert "| Contacts | 56 |" in report
    assert "| Contacts using the real inbox | 3 |" in report
    assert "| Branches not covered | 1 |" in report
    # rows per CSV match the files
    for name in result.data.tables:
        _, rows = read_csv(result.files[table_file_name(name)])
        assert f"| {table_file_name(name)} | {len(rows)} |" in report
    assert "| DE_Log | written | - | no | - |" in report


def test_report_lists_checks_branches_unsupported_warnings_and_settings():
    report = write_report(build_result())
    assert "1 of the 2 trust checks passed. 1 failed: the details are below." in report
    assert "| Every key joins | FAILED | 2 |" in report
    assert "- DE_Master row 4 has no scenario row." in report
    assert "| Clicked Promo (ENG-2) | Yes (yes) | The email Promo has no link, so a click cannot be asked. |" in report
    assert "`C.Region LIKE 'N%' OR C.Region IS NULL`" in report
    assert "Placeholder: DE_Entry.Weird" in report and "Placeholder: DE_Entry.Mandatory" in report
    for warning in ("Reading warning from the loader.", "Table selection warning.",
                    "Parser warning about the entry query.", "Coverage warning."):
        assert f"- {warning}" in report
    assert "| Real inbox for open and click tests | tester@example.com |" in report
    assert "| One active scenario at a time, as applied | on, activation field DE_Config.ACTIVE |" in report
    assert "| Generation time | 2026-10-11 09:30 |" in report
    assert "W. Europe Standard Time" in report


def test_report_scenarios_and_cases():
    report = write_report(build_result())
    assert ("| TST_R01_S01 | 1 | yes | True | False | 2 | 6 | TST_R01_C0001 to TST_R01_C0006 | "
            "DE_Config.CHANNEL = Email |") in report
    assert "| S02-C01 | TST_R01_S02 | 50 | TST_R01_C0007 to TST_R01_C0056 |" in report
    assert "1 to 9 contacts expected on the 10% side" in report


def test_report_without_checks_and_without_problems():
    result = build_result(checks=[])
    result.coverage.branches = [b for b in result.coverage.branches if b.status != "not covered"]
    report = write_report(result)
    assert "No trust check was run on these files." in report
    assert "Every branch is covered." in report


def test_suspicious_patterns_are_gathered_from_every_module():
    result = build_result()
    result.coverage.warnings.append("suspicious pattern: DE_Config.CHANNEL is compared with 'Email' and 'EMAIL'.")
    result.package.warnings.append("Suspicious pattern: two splits share the name 'Channel'.")
    report = write_report(result)
    section = report[report.index("### Suspicious patterns"):report.index("### While reading the export")]
    assert "- DE_Config.CHANNEL is compared with 'Email' and 'EMAIL'." in section
    assert "- Two splits share the name 'Channel'." in section
    assert "suspicious pattern:" not in report.lower().replace("suspicious patterns", "")


# ---------------------------------------------------------------------------
# Testbook
# ---------------------------------------------------------------------------

def test_testbook_has_one_row_per_case_and_names_every_contact():
    result = build_result()
    rows = build_testbook_rows(result)
    assert len(rows) == len(result.coverage.cases) == 3
    assert [row[0] for row in rows] == [1, 2, 3]
    assert {row[1] for row in rows} == {"SFMC"} and {row[7] for row in rows} == {"TO START"}
    assert [row[2] for row in rows] == ["S01-C01 - Welcome opened", "S01-C02 - Path ending at Wait one minute",
                                        "S02-C01 - Random split 90/10"]
    text = "\n".join(str(value) for row in rows for value in row)
    for contact in result.data.contacts:
        if contact.case_id != "S02-C01":
            assert contact.key in text
    assert "Test contacts (50): TST_R01_C0007 to TST_R01_C0056." in rows[2][4]
    random_keys = [c.key for c in result.data.contacts if c.case_id == "S02-C01"]
    assert random_keys == [f"TST_R01_C{n:04d}" for n in range(7, 57)]   # the range names exactly these


def test_testbook_preconditions_steps_and_expected_result():
    rows = build_testbook_rows(build_result())
    first = dict(zip(TESTBOOK_COLUMNS, rows[0]))
    assert "Import these CSV files: DE_Entry.csv, DE_Master.csv, DE_Config.csv." in first["Pre conditions"]
    assert "Active scenario: TST_R01_S01" in first["Pre conditions"]
    assert "Expected path: Channel [Email] > Welcome > Wait campaign date > Wait reminder date > " \
           "Opened Welcome [Yes] > Log open." in first["Description"]
    steps = first["Detailed steps"].split("\n")
    assert [step.split(".")[0] for step in steps] == ["1", "2", "3", "4", "5"]
    assert "Before 2026-10-12 09:30, open the email 'Welcome'" in steps[3]
    assert "tester+TST_R01_C0001@example.com" in steps[3]
    assert "- Emails received: Welcome." in first["Expected result"]
    assert "- Values written by the journey: DE_Log.Status = Opened." in first["Expected result"]
    assert "- Last activity reached: Log open (UPD-1)." in first["Expected result"]
    assert first["Note"].startswith("Manual engagement required: open the email 'Welcome' before 2026-10-12 09:30.")

    second = dict(zip(TESTBOOK_COLUMNS, rows[1]))
    assert "Contact profile: DE_Master.Perm = True." in second["Description"]     # CSV form of "true"
    assert "Wait one minute (1 minute)" in second["Detailed steps"]
    assert second["Note"] == ""

    third = dict(zip(TESTBOOK_COLUMNS, rows[2]))
    assert "activate scenario TST_R01_S02: in DE_Config, set ACTIVE to True on the row with CampaignName " \
           "TST_R01_S02 and to False on every other row." in third["Detailed steps"]
    assert "- 10%: emails received: Promo; values written: none; last activity reached: Clicked Promo (ENG-2)." \
           in third["Expected result"]
    assert "Accepted range: 1 to 9 contacts on the 10% side." in third["Expected result"]
    assert "Between 1 and 9 of the 50 contacts are expected on the 10% side." in third["Note"]


def test_testbook_without_one_active_scenario_or_real_inbox():
    rows = build_testbook_rows(build_result(one_active_scenario=False, real_inbox=""))
    first = dict(zip(TESTBOOK_COLUMNS, rows[0]))
    assert "Every scenario is active at the same time." in first["Pre conditions"]
    assert "no real inbox was set" in first["Detailed steps"]


def test_testbook_xlsx_layout():
    result = build_result()
    sheet = load_workbook(io.BytesIO(result.files[TESTBOOK_FILE])).active
    values = list(sheet.iter_rows(values_only=True))
    assert list(values[0]) == TESTBOOK_COLUMNS
    assert len(values) == 1 + len(result.coverage.cases)
    assert [list(row) for row in values[1:]] == [[v if v != "" else None for v in row] for row in build_testbook_rows(result)]


# ---------------------------------------------------------------------------
# Determinism
# ---------------------------------------------------------------------------

def test_two_runs_give_identical_bytes():
    first, second = build_result(), build_result()
    assert first.files == second.files
    assert build_zip(first.files) == build_zip(second.files)
    assert write_testbook(first) == write_testbook(first)
